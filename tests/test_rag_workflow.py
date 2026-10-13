import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import docx
import openpyxl
from fastapi.testclient import TestClient

import chatbot


class FakeEmbedder:
    def encode(self, texts, **_kwargs):
        return np.tile(np.array([[1.0, 0.0]], dtype=np.float32), (len(texts), 1))


class RagWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        upload_dir = Path(self.temp.name) / "uploads"
        upload_dir.mkdir()
        fake_embedder = FakeEmbedder()
        for name, value in (
            ("INDEX_DIR", Path(self.temp.name)),
            ("UPLOADS_DIR", upload_dir),
            ("get_embedder", lambda: fake_embedder),
            ("_rag_engine", None),
            ("_rag_load_error", None),
            ("APP_ACCESS_TOKEN", ""),
            ("HOST", "127.0.0.1"),
        ):
            patcher = patch.object(chatbot, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    @staticmethod
    def upload(client, name, text):
        return client.post("/upload", files=[("files", (name, text.encode("utf-8")))])

    def test_upload_appends_and_same_name_replaces(self):
        with TestClient(chatbot.app) as client:
            self.assertEqual(self.upload(client, "A.txt", "first A").status_code, 200)
            self.assertEqual(self.upload(client, "B.txt", "first B").status_code, 200)
            self.assertEqual(
                [doc["source"] for doc in chatbot._rag_engine["docs"]],
                ["A.txt", "B.txt"],
            )
            self.assertEqual(self.upload(client, "A.txt", "second A").status_code, 200)
            self.assertEqual(
                [(doc["source"], doc["text"]) for doc in chatbot._rag_engine["docs"]],
                [("B.txt", "first B"), ("A.txt", "second A")],
            )
            self.assertEqual(client.get("/status").json()["rag_chunks"], 2)
            self.assertTrue((Path(self.temp.name) / "index.npz").exists())
            self.assertEqual((Path(self.temp.name) / "uploads" / "A.txt").read_text(), "second A")
            self.assertEqual((Path(self.temp.name) / "uploads" / "B.txt").read_text(), "first B")
            self.assertEqual(
                [(entry["name"], entry["stored"]) for entry in client.get("/files").json()["files"]],
                [("A.txt", True), ("B.txt", True)],
            )

    def test_index_reloads_and_reset_removes_saved_files(self):
        with TestClient(chatbot.app) as client:
            self.assertEqual(self.upload(client, "A.txt", "saved text").status_code, 200)
            chatbot._rag_engine = None
            self.assertEqual(chatbot.load_index()["docs"][0]["text"], "saved text")
            self.assertEqual(client.post("/reset_index").status_code, 200)
            self.assertEqual(client.get("/status").json()["rag_chunks"], 0)
            self.assertFalse((Path(self.temp.name) / "index.npz").exists())
            self.assertEqual(client.get("/files").json()["files"], [])
            self.assertFalse((Path(self.temp.name) / "uploads" / "A.txt").exists())

    def test_legacy_index_loads_without_deleting_user_data(self):
        root = Path(self.temp.name)
        (root / "docs.json").write_text(
            json.dumps([{"source": "old.txt", "text": "legacy"}]), encoding="utf-8"
        )
        np.save(root / "embeddings.npy", np.array([[1.0, 0.0]], dtype=np.float32))
        self.assertEqual(chatbot.load_index()["docs"][0]["source"], "old.txt")
        self.assertEqual(chatbot.list_uploaded_files()[0]["stored"], False)
        self.assertTrue((root / "docs.json").exists())
        self.assertTrue((root / "embeddings.npy").exists())

    def test_corrupt_index_is_reported_and_not_overwritten(self):
        archive = Path(self.temp.name) / "index.npz"
        archive.write_bytes(b"not an npz archive")
        with TestClient(chatbot.app) as client:
            self.assertIsNotNone(client.get("/status").json()["rag_error"])
            self.assertEqual(self.upload(client, "A.txt", "new data").status_code, 400)
            self.assertEqual(archive.read_bytes(), b"not an npz archive")

    def test_upload_limits_and_duplicate_names(self):
        with TestClient(chatbot.app) as client, patch.object(chatbot, "MAX_UPLOAD_BYTES", 4):
            self.assertEqual(self.upload(client, "large.txt", "12345").status_code, 413)
            response = client.post("/upload", files=[
                ("files", ("x.txt", b"one")), ("files", ("x.txt", b"two")),
            ])
            self.assertEqual(response.status_code, 400)
            self.assertEqual(client.get("/status").json()["rag_chunks"], 0)

    def test_batch_and_total_upload_limits(self):
        with TestClient(chatbot.app) as client:
            too_many = [("files", (f"{i}.txt", b"x")) for i in range(11)]
            self.assertEqual(client.post("/upload", files=too_many).status_code, 413)
            with patch.object(chatbot, "MAX_TOTAL_UPLOAD_BYTES", 5):
                files = [("files", ("a.txt", b"abc")), ("files", ("b.txt", b"def"))]
                self.assertEqual(client.post("/upload", files=files).status_code, 413)
            self.assertEqual(client.get("/status").json()["rag_chunks"], 0)

    def test_docx_tables_and_all_excel_sheets_are_extracted(self):
        document = docx.Document()
        document.add_paragraph("正文")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "名称"
        table.cell(0, 1).text = "数值"
        document_bytes = io.BytesIO()
        document.save(document_bytes)
        parsed_doc = chatbot.extract_text("sample.docx", document_bytes.getvalue())
        self.assertIn("正文", parsed_doc)
        self.assertIn("名称 | 数值", parsed_doc)

        workbook = openpyxl.Workbook()
        workbook.active.title = "第一表"
        workbook.active.append(["名称", "甲"])
        second = workbook.create_sheet("第二表")
        second.append(["名称", "乙"])
        workbook_bytes = io.BytesIO()
        workbook.save(workbook_bytes)
        parsed_book = chatbot.extract_text("sample.xlsx", workbook_bytes.getvalue())
        self.assertIn("[工作表: 第一表]", parsed_book)
        self.assertIn("[工作表: 第二表]", parsed_book)

    def test_token_protects_data_routes_and_guide_is_public(self):
        with patch.object(chatbot, "APP_ACCESS_TOKEN", "test-access-token"):
            with TestClient(chatbot.app) as client:
                self.assertEqual(client.get("/status").status_code, 401)
                self.assertEqual(client.post("/reset_index").status_code, 401)
                self.assertEqual(client.post("/chat", json={"message": "hello"}).status_code, 401)
                self.assertEqual(client.get("/files").status_code, 401)
                self.assertEqual(client.delete("/files", params={"name": "A.txt"}).status_code, 401)
                self.assertEqual(client.get("/guide").status_code, 200)
                response = client.get(
                    "/status", headers={"Authorization": "Bearer test-access-token"}
                )
                self.assertEqual(response.status_code, 200)
                self.assertNotIn("test-access-token", response.text)

    def test_delete_one_file_updates_directory_index_and_restart(self):
        with TestClient(chatbot.app) as client:
            self.assertEqual(self.upload(client, "A.txt", "first A").status_code, 200)
            self.assertEqual(self.upload(client, "B.txt", "first B").status_code, 200)
            response = client.delete("/files", params={"name": "A.txt"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["chunks"], 1)
            self.assertEqual([x["name"] for x in client.get("/files").json()["files"]], ["B.txt"])
            self.assertFalse((Path(self.temp.name) / "uploads" / "A.txt").exists())
            chatbot._rag_engine = None
            self.assertEqual(chatbot.load_index()["docs"][0]["source"], "B.txt")
            self.assertEqual(client.delete("/files", params={"name": "A.txt"}).status_code, 404)

    def test_delete_last_file_keeps_valid_empty_index(self):
        with TestClient(chatbot.app) as client:
            self.assertEqual(self.upload(client, "A.txt", "only A").status_code, 200)
            self.assertEqual(client.delete("/files", params={"name": "A.txt"}).status_code, 200)
            self.assertEqual(client.get("/status").json()["rag_chunks"], 0)
            chatbot._rag_engine = None
            restored = chatbot.load_index()
            self.assertEqual(restored["docs"], [])
            self.assertEqual(client.get("/files").json()["files"], [])

    def test_index_save_failure_restores_original_file(self):
        with TestClient(chatbot.app) as client:
            self.assertEqual(self.upload(client, "A.txt", "original").status_code, 200)
            original = (Path(self.temp.name) / "uploads" / "A.txt")
            with patch.object(chatbot, "save_index", side_effect=OSError("disk full")):
                with self.assertRaises(OSError):
                    chatbot.build_index([{"name": "A.txt", "data": b"replacement"}])
                with self.assertRaises(OSError):
                    chatbot.remove_uploaded_file("A.txt")
            self.assertEqual(original.read_bytes(), b"original")
            self.assertEqual(chatbot._rag_engine["docs"][0]["text"], "original")

    def test_rejects_unsafe_upload_name(self):
        with TestClient(chatbot.app) as client:
            response = client.post("/upload", files=[("files", ("bad:name.txt", b"data"))])
            self.assertEqual(response.status_code, 400)
            self.assertEqual(client.get("/files").json()["files"], [])

    def test_remote_bind_requires_access_token(self):
        with patch.object(chatbot, "HOST", "0.0.0.0"):
            with self.assertRaisesRegex(ValueError, "APP_ACCESS_TOKEN"):
                with TestClient(chatbot.app):
                    pass
            with patch.object(chatbot, "APP_ACCESS_TOKEN", "too-short"):
                with self.assertRaisesRegex(ValueError, "16"):
                    with TestClient(chatbot.app):
                        pass


if __name__ == "__main__":
    unittest.main()
