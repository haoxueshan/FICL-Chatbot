"""
FICL Chatbot — RAG + LLM 챗봇 학습 예제
─────────────────────────────────────────
- GGUF 양자화 모델을 llama-cpp-python으로 로컬 실행
- 파일 업로드 → 임베딩 → 의미 검색 (RAG)
- FastAPI + SSE(Server-Sent Events)로 토큰 단위 스트리밍

학습 목표:
1. GGUF 양자화 모델을 일반 PC에서 어떻게 돌리는가
2. RAG(검색 증강 생성)가 무엇이며 어떻게 구현되는가
3. LLM 응답을 어떻게 토큰 단위로 스트리밍하는가
"""

import os, io, re, json, hmac, tempfile, threading, uuid
from contextlib import asynccontextmanager
from pathlib import Path
from typing import List, Dict, Literal, Optional
from urllib.parse import urlparse

from fastapi import FastAPI, File, UploadFile, HTTPException, Depends, Request, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, HTMLResponse
from pydantic import BaseModel, ConfigDict, Field
from starlette.concurrency import run_in_threadpool

import numpy as np
import httpx

# ─────────────────────────────────────────────
# 외부 라이브러리
# ─────────────────────────────────────────────
# LLM 로딩 (GGUF 모델)
try:
    from llama_cpp import Llama
    LLAMA_OK = True
except ImportError:
    LLAMA_OK = False
    print("[경고] llama-cpp-python 미설치. README.md를 참고해 설치하세요.")

# 임베딩 (RAG용)
try:
    from sentence_transformers import SentenceTransformer
    ST_OK = True
except ImportError:
    ST_OK = False

# 문서 파서
try:
    import pypdf
except ImportError:
    pypdf = None
try:
    import docx
except ImportError:
    docx = None
try:
    import pandas as pd
except ImportError:
    pd = None


# ═══════════════════════════════════════════════════════
# 설정
# ═══════════════════════════════════════════════════════
BASE_DIR = Path(__file__).parent
MODELS_DIR = BASE_DIR / "models"
INDEX_DIR = BASE_DIR / "rag_index"
UPLOADS_DIR = BASE_DIR / "uploads"
UI_HTML = BASE_DIR / "ui.html"
GUIDE_HTML = BASE_DIR / "guide.html"

MODELS_DIR.mkdir(exist_ok=True)
INDEX_DIR.mkdir(exist_ok=True)
UPLOADS_DIR.mkdir(exist_ok=True)

# 지원 모델 카탈로그 — 학생이 직접 다운로드해서 파일명대로 models/에 넣음
MODELS = {
    "qwen2.5-1.5b": {
        "file": "qwen2.5-1.5b-instruct-q4_k_m.gguf",
        "name": "Qwen2.5-1.5B Q4_K_M（测试用 · 约 1GB）",
        "context": 8192,
        "note": "轻量测试模型，CPU 也可运行",
    },
    "qwen3-8b": {
        "file": "Qwen3-8B-Q4_K_M.gguf",
        "name": "Qwen3-8B Q4_K_M",
        "context": 8192,
        "note": "多语言均衡型",
    },
    "eeve-10.8b": {
        "file": "EEVE-Korean-Instruct-10.8B-v1.0-Q4_K_M.gguf",
        "name": "EEVE-Korean-10.8B Q4_K_M（韩语优化）",
        "context": 4096,
        "note": "韩语优化",
    },
    "exaone-7.8b": {
        "file": "EXAONE-3.5-7.8B-Instruct-Q4_K_M.gguf",
        "name": "EXAONE-3.5-7.8B Q4_K_M（韩语优化）",
        "context": 4096,
        "note": "LG AI，韩语优化",
    },
}

DEFAULT_MODEL = os.environ.get("DEFAULT_MODEL", "qwen3-8b")
HOST = os.environ.get("HOST", "127.0.0.1")
PORT = int(os.environ.get("PORT", "7861"))
APP_ACCESS_TOKEN = os.environ.get("APP_ACCESS_TOKEN", "").strip()
N_GPU_LAYERS = int(os.environ.get("N_GPU_LAYERS", "-1"))  # -1: 전체 GPU, 0: CPU만

# OpenAI 호환 API (선택). 모델 및 키는 서버 전용 파일에서 읽는다.
API_MODELS_FILE = BASE_DIR / "api_models.json"

EMBED_MODEL_NAME = os.environ.get("EMBED_MODEL", "intfloat/multilingual-e5-base")
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_UPLOAD_FILES = 10
MAX_TOTAL_UPLOAD_BYTES = 50 * 1024 * 1024
ALLOWED_EXT = {".pdf", ".docx", ".xlsx", ".txt", ".md"}


# ═══════════════════════════════════════════════════════
# LLM 로딩 (지연 로딩 + 락 보호)
# ═══════════════════════════════════════════════════════
_llm = None
_llm_key = None
_llm_lock = threading.Lock()


# 키워드 매핑 — 파일명이 정확히 일치하지 않아도 키워드로 자동 인식
_MODEL_KEYWORDS = {
    "qwen2.5-1.5b":["qwen2.5-1.5b", "qwen2_5-1_5b", "qwen-2.5-1.5b"],
    "qwen3-8b":    ["qwen3-8b", "qwen3_8b", "qwen-3-8b"],
    "eeve-10.8b":  ["eeve", "eeve-korean"],
    "exaone-7.8b": ["exaone"],
}


def find_gguf_file(model_key: str) -> Optional[Path]:
    """모델 키에 해당하는 GGUF 파일 찾기.
    1) MODELS의 정확한 파일명으로 우선 매칭
    2) 안 되면 _MODEL_KEYWORDS의 키워드로 fuzzy 매칭 (대소문자 무시)
    3) 그래도 없으면 None
    """
    if model_key not in MODELS:
        return None
    # 1) 정확한 파일명
    exact = MODELS_DIR / MODELS[model_key]["file"]
    if exact.exists():
        return exact
    # 2) 키워드 fuzzy 매칭 (.gguf 파일 중 키워드 포함)
    keywords = _MODEL_KEYWORDS.get(model_key, [model_key])
    for f in sorted(MODELS_DIR.glob("*.gguf")):
        name_lower = f.name.lower()
        if any(kw.lower() in name_lower for kw in keywords):
            return f
    return None


def get_llm(model_key: str):
    """요청한 로컬 모델만 로딩한다. 선택한 모델을 다른 모델로 바꾸지 않는다."""
    global _llm, _llm_key
    if not LLAMA_OK:
        raise RuntimeError("未安装 llama-cpp-python；请参阅 README.md。")

    if model_key not in MODELS:
        raise ValueError(f"未注册的模型：{model_key}")

    mp = find_gguf_file(model_key)
    if mp is None:
        raise FileNotFoundError(f"找不到所选模型文件：{model_key}")

    # fast path — 이미 같은 모델 로딩됨
    if _llm is not None and _llm_key == model_key:
        return _llm

    with _llm_lock:
        if _llm is not None and _llm_key == model_key:
            return _llm
        # 기존 모델 해제 (메모리 정리)
        if _llm is not None:
            del _llm
            _llm = None

        print(f"[LLM 로딩 시작] {MODELS[model_key]['name']}")
        print(f"  파일: {mp}")
        print(f"  GPU 레이어: {N_GPU_LAYERS} (-1=전체, 0=CPU)")

        _llm = Llama(
            model_path=str(mp),
            n_ctx=MODELS[model_key]["context"],
            n_gpu_layers=N_GPU_LAYERS,
            verbose=False,
        )
        _llm_key = model_key
        print(f"[LLM 로딩 완료] {MODELS[model_key]['name']}\n")
        return _llm


def build_api_chat_url(base_url: str) -> str:
    """OpenAI 호환 Chat Completions URL 검증·정규화.

    외부 호스트는 HTTPS만 허용한다. 개발용 로컬 서버는 HTTP를 허용한다.
    """
    base = base_url.strip().rstrip("/")
    parsed = urlparse(base)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ValueError("API Base URL은 http(s)://로 시작하는 올바른 주소여야 합니다.")
    if parsed.username or parsed.password or parsed.query or parsed.fragment:
        raise ValueError("API Base URL에는 계정 정보, 쿼리 또는 fragment를 넣을 수 없습니다.")
    is_loopback = parsed.hostname.lower() in ("localhost", "127.0.0.1", "::1")
    if parsed.scheme != "https" and not is_loopback:
        raise ValueError("외부 API는 HTTPS만 허용됩니다. HTTP는 localhost에서만 사용할 수 있습니다.")
    if base.endswith("/chat/completions"):
        return base
    if base.endswith("/v1"):
        return base + "/chat/completions"
    return base + "/v1/chat/completions"


def get_api_models() -> Dict[str, Dict]:
    """서버 전용 api_models.json에 등록된 모델만 선택 가능하게 한다."""
    providers = []
    if API_MODELS_FILE.exists():
        extra = json.loads(API_MODELS_FILE.read_text(encoding="utf-8"))
        if not isinstance(extra, list):
            raise ValueError("api_models.json은 provider 목록(JSON 배열)이어야 합니다.")
        providers.extend(extra)

    catalog = {}
    for provider in providers:
        if not isinstance(provider, dict):
            raise ValueError("API provider 항목은 JSON 객체여야 합니다.")
        provider_id = provider.get("id")
        if not isinstance(provider_id, str) or not re.fullmatch(r"[A-Za-z0-9_-]+", provider_id):
            raise ValueError("API provider id는 영숫자, _ 또는 -만 사용할 수 있습니다.")
        base_url = provider.get("base_url")
        if not isinstance(base_url, str):
            raise ValueError("API provider base_url이 필요합니다.")
        build_api_chat_url(base_url)
        configured_key = provider.get("api_key", "")
        if not isinstance(configured_key, str):
            raise ValueError(f"Provider {provider_id}: api_key 必须是字符串。")
        configured_key = configured_key.strip()
        model_list = provider.get("models")
        if not isinstance(model_list, list):
            raise ValueError("API provider의 models는 배열이어야 합니다.")
        for item in model_list:
            if not isinstance(item, (str, dict)):
                raise ValueError("API models 항목은 문자열 또는 JSON 객체여야 합니다.")
            model_id = item if isinstance(item, str) else item.get("id")
            name = model_id if isinstance(item, str) else item.get("name", model_id)
            if not isinstance(model_id, str) or not model_id.strip() or len(model_id) > 200:
                raise ValueError("API 모델 id가 올바르지 않습니다.")
            model_id = model_id.strip()
            enable_thinking = item.get("enable_thinking") if isinstance(item, dict) else None
            if enable_thinking is not None and not isinstance(enable_thinking, bool):
                raise ValueError("enable_thinking은 true 또는 false여야 합니다.")
            key = f"api:{provider_id}:{model_id}"
            if key in catalog:
                raise ValueError(f"중복 API 모델: {key}")
            catalog[key] = {
                "key": key,
                "name": f"{provider.get('name', provider_id)} · {name}",
                "provider": "api",
                "base_url": base_url,
                "model_id": model_id,
                "api_key": configured_key,
                "enable_thinking": enable_thinking,
                "available": bool(configured_key),
            }
    return catalog


def get_model_catalog() -> Dict[str, Dict]:
    """로컬/API 모델을 동일한 선택 키와 상태 형식으로 모은다."""
    catalog = {}
    for key, spec in MODELS.items():
        found = find_gguf_file(key)
        catalog[key] = {
            "key": key,
            "name": spec["name"],
            "provider": "local",
            "note": spec["note"],
            "downloaded": found is not None,
            "file": found.name if found else None,
            "available": found is not None and LLAMA_OK,
            "unavailable_reason": None if found and LLAMA_OK else (
                "未安装 llama-cpp-python" if found else "尚未下载模型文件"
            ),
        }
    for key, spec in get_api_models().items():
        catalog[key] = {
            **spec,
            "note": "兼容 OpenAI 的 API",
            "downloaded": False,
            "file": None,
            "unavailable_reason": None if spec["available"] else "后端未配置 API Key",
        }
    return catalog


def get_default_model_key(catalog: Dict[str, Dict]) -> Optional[str]:
    """설정된 기본 모델이 사용 가능하면 우선하고, 아니면 첫 가용 모델을 고른다."""
    if DEFAULT_MODEL in catalog and catalog[DEFAULT_MODEL]["available"]:
        return DEFAULT_MODEL
    return next((key for key, model in catalog.items() if model["available"]), None)


def select_model(requested_key: Optional[str], catalog: Dict[str, Dict]) -> Dict:
    """상태 조회와 채팅에 사용한 동일한 카탈로그로 모델 선택을 검증한다."""
    key = requested_key if requested_key is not None else get_default_model_key(catalog)
    if key is None:
        raise HTTPException(503, "没有可用模型，请检查 /status。")
    model = catalog.get(key)
    if model is None:
        raise HTTPException(400, "模型未注册，请检查 /status 中的模型列表。")
    if not model["available"]:
        raise HTTPException(400, f"所选模型不可用：{model['unavailable_reason']}")
    return model


def public_model_info(model: Dict) -> Dict:
    """브라우저에는 선택에 필요한 정보만 전달하고 서버 설정·키는 숨긴다."""
    return {field: model[field] for field in (
        "key", "name", "provider", "note", "downloaded", "file",
        "available", "unavailable_reason",
    )}


def stream_api_completion(messages: List[Dict], selected_model: Dict):
    """OpenAI 호환 API 응답을 토큰 문자열 iterator로 변환."""
    key = selected_model.get("api_key", "")
    model = selected_model["model_id"]
    if not key:
        raise ValueError("所选 API 模型未在 api_models.json 中配置 api_key。")

    url = build_api_chat_url(selected_model["base_url"])
    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    payload = {
        "model": model,
        "messages": messages,
        "stream": True,
    }
    if selected_model.get("enable_thinking") is not None:
        payload["enable_thinking"] = selected_model["enable_thinking"]

    timeout = httpx.Timeout(120.0, connect=15.0)
    with httpx.Client(timeout=timeout, follow_redirects=False) as client:
        with client.stream("POST", url, headers=headers, json=payload) as response:
            if response.status_code == 401:
                raise RuntimeError(
                    f"API 认证失败 (401)：{selected_model['name']} 的密钥无效，"
                    "请检查 API Key 与接口地址、服务及区域是否匹配。"
                )
            if response.status_code >= 400:
                detail = response.read().decode("utf-8", errors="replace")[:500]
                detail = detail.replace(key, "***")
                raise RuntimeError(f"API 요청 실패 ({response.status_code}): {detail}")

            content_type = response.headers.get("content-type", "").lower()
            if "text/event-stream" not in content_type:
                data = json.loads(response.read())
                content = data.get("choices", [{}])[0].get("message", {}).get("content", "")
                if content:
                    yield content
                return

            for line in response.iter_lines():
                if not line or not line.startswith("data:"):
                    continue
                raw = line[5:].strip()
                if raw == "[DONE]":
                    break
                try:
                    event = json.loads(raw)
                    if event.get("error"):
                        detail = json.dumps(event["error"], ensure_ascii=False)[:500]
                        detail = detail.replace(key, "***")
                        raise RuntimeError(f"API 스트림 오류: {detail}")
                    delta = event.get("choices", [{}])[0].get("delta", {}).get("content", "")
                except RuntimeError:
                    raise
                except (json.JSONDecodeError, KeyError, IndexError, AttributeError):
                    delta = ""
                if isinstance(delta, str) and delta:
                    yield delta


def verify_api_model(selected_model: Dict) -> bool:
    """명시적으로 요청받았을 때만 아주 짧은 실 API 호출로 인증을 확인한다."""
    probe_model = dict(selected_model)
    if probe_model.get("enable_thinking") is not None:
        probe_model["enable_thinking"] = False
    stream = stream_api_completion(
        [{"role": "user", "content": "请只回复 OK。"}], probe_model
    )
    try:
        if not next(stream, None):
            raise RuntimeError("API 已连接，但未返回可用的回答内容。")
        return True
    finally:
        close = getattr(stream, "close", None)
        if close:
            close()

# ═══════════════════════════════════════════════════════
# RAG (임베딩 + 코사인 유사도 검색)
# ═══════════════════════════════════════════════════════
_embedder = None
_embedder_lock = threading.Lock()
_rag_engine: Optional[Dict] = None   # {"docs": [...], "embeddings": np.ndarray}
_rag_lock = threading.RLock()
_rag_load_error: Optional[str] = None


def get_embedder():
    """임베딩 모델 (한 번만 로딩)."""
    global _embedder
    if _embedder is None:
        with _embedder_lock:
            if _embedder is None:
                if not ST_OK:
                    raise RuntimeError("sentence-transformers 미설치")
                print(f"[임베딩 모델 로딩] {EMBED_MODEL_NAME}")
                _embedder = SentenceTransformer(EMBED_MODEL_NAME)
    return _embedder


def chunk_text(text: str, size: int = 512, overlap: int = 128) -> List[str]:
    """긴 문서를 검색 단위로 자르기."""
    text = re.sub(r"\s+", " ", text or "").strip()
    if not text:
        return []
    chunks, start = [], 0
    while start < len(text):
        end = min(len(text), start + size)
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start = end - overlap
    return chunks


def extract_text(name: str, data: bytes) -> str:
    """파일 확장자별로 텍스트 추출."""
    ext = os.path.splitext(name)[1].lower()
    if ext == ".pdf":
        if not pypdf:
            raise RuntimeError("缺少 PDF 解析依赖 pypdf。")
        reader = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join(p.extract_text() or "" for p in reader.pages)
    if ext == ".docx":
        if not docx:
            raise RuntimeError("缺少 DOCX 解析依赖 python-docx。")
        d = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in d.paragraphs]
        tables = [" | ".join(cell.text for cell in row.cells)
                  for table in d.tables for row in table.rows]
        return "\n".join(paragraphs + tables)
    if ext == ".xlsx":
        if pd is None:
            raise RuntimeError("缺少 XLSX 解析依赖 pandas/openpyxl。")
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        return "\n".join(
            f"[工作表: {sheet}]\n{frame.to_csv(index=False)}"
            for sheet, frame in sheets.items()
        )
    if ext in (".txt", ".md"):
        for encoding in ("utf-8-sig", "gb18030", "cp949"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                pass
        raise ValueError(f"无法识别文本编码：{name}；请转为 UTF-8。")
    return ""


def validate_upload_name(name: str) -> str:
    """업로드 디렉터리 안의 단일 파일명만 허용한다."""
    if (not isinstance(name, str) or not name or len(name) > 200
            or name in (".", "..") or name != name.strip(" .")
            or any(c in name for c in '/\\<>:"|?*')
            or any(ord(c) < 32 for c in name)
            or name.split(".")[0].upper() in {
                "CON", "PRN", "AUX", "NUL", "COM1", "COM2", "COM3", "COM4",
                "COM5", "COM6", "COM7", "COM8", "COM9", "LPT1", "LPT2",
                "LPT3", "LPT4", "LPT5", "LPT6", "LPT7", "LPT8", "LPT9",
            }):
        raise ValueError("文件名不安全或不受支持。")
    if os.path.splitext(name)[1].lower() not in ALLOWED_EXT:
        raise ValueError(f"不支持的文件格式：{name}")
    return name


def list_uploaded_files() -> List[Dict]:
    """디스크의 원본 파일과 기존 인덱스만 남은 자료를 함께 표시한다."""
    with _rag_lock:
        chunk_counts = {}
        if _rag_engine:
            for doc in _rag_engine["docs"]:
                name = doc["source"]
                chunk_counts[name] = chunk_counts.get(name, 0) + 1
        names = set(chunk_counts)
        names.update(p.name for p in UPLOADS_DIR.iterdir() if p.is_file()
                     and not p.name.startswith("."))
        result = []
        for name in sorted(names, key=str.casefold):
            path = UPLOADS_DIR / name
            stored = path.parent == UPLOADS_DIR and path.is_file()
            result.append({
                "name": name,
                "size": path.stat().st_size if stored else None,
                "chunks": chunk_counts.get(name, 0),
                "stored": stored,
            })
        return result


def build_index(files: List[Dict]) -> Optional[Dict]:
    """새 파일을 추가하고, 같은 이름의 기존 파일만 교체한다."""
    global _rag_engine
    if _rag_load_error:
        raise RuntimeError(f"现有索引无法加载：{_rag_load_error}。请先备份或清空索引。")
    embedder = get_embedder()

    all_chunks = []
    for f in files:
        try:
            text = extract_text(f["name"], f["data"])
        except (ValueError, RuntimeError):
            raise
        except Exception as e:
            raise ValueError(f"无法读取文件 {f['name']}：{type(e).__name__}") from e
        for chunk in chunk_text(text):
            all_chunks.append({"source": f["name"], "text": chunk})

    if not all_chunks:
        return None

    print(f"[RAG] {len(all_chunks)}개 청크 임베딩 중...")
    # E5 계열은 문서(passage)에 "passage: " 프리픽스 필수 (검색 품질 30%+ 향상)
    # 다른 모델로 교체해도 프리픽스는 무해함
    texts = [f"passage: {c['text']}" for c in all_chunks]
    embeddings = embedder.encode(texts, normalize_embeddings=True, show_progress_bar=False)

    new_embeddings = np.asarray(embeddings, dtype=np.float32)
    if new_embeddings.ndim != 2 or len(new_embeddings) != len(all_chunks):
        raise RuntimeError("嵌入模型返回的数据数量或维度不正确。")

    UPLOADS_DIR.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(dir=UPLOADS_DIR, prefix=".pending-") as pending:
        staged = []
        for i, file in enumerate(files):
            temporary = Path(pending) / f"{i}.new"
            temporary.write_bytes(file["data"])
            staged.append(temporary)

        with _rag_lock:
            if _rag_engine:
                replacing = {f["name"].casefold() for f in files}
                keep = [i for i, doc in enumerate(_rag_engine["docs"])
                        if doc["source"].casefold() not in replacing]
                old_embeddings = _rag_engine["embeddings"]
                if old_embeddings.shape[1] != new_embeddings.shape[1]:
                    raise RuntimeError("嵌入维度已变化，请先备份并清空旧索引。")
                docs = [_rag_engine["docs"][i] for i in keep] + all_chunks
                merged = np.concatenate((old_embeddings[keep], new_embeddings), axis=0)
            else:
                docs, merged = all_chunks, new_embeddings

            engine = {"docs": docs, "embeddings": merged}
            promoted = []
            try:
                for i, file in enumerate(files):
                    destination = UPLOADS_DIR / file["name"]
                    backup = Path(pending) / f"{i}.old" if destination.exists() else None
                    if backup:
                        os.replace(destination, backup)
                    promoted.append((destination, backup))
                    os.replace(staged[i], destination)
                save_index(engine)
            except Exception:
                for destination, backup in reversed(promoted):
                    if destination.exists():
                        destination.unlink()
                    if backup and backup.exists():
                        os.replace(backup, destination)
                raise
            _rag_engine = engine
            print(f"[RAG] 인덱스 저장 완료: {len(docs)}개 청크")
            return engine


def save_index(engine: Dict) -> None:
    """단일 파일로 저장하여 문서와 벡터가 서로 다른 버전이 되는 일을 방지한다."""
    payload = json.dumps(engine["docs"], ensure_ascii=False).encode("utf-8")
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(dir=INDEX_DIR, suffix=".npz", delete=False) as tmp:
            temp_path = Path(tmp.name)
            np.savez_compressed(
                tmp,
                embeddings=engine["embeddings"],
                docs=np.frombuffer(payload, dtype=np.uint8),
                embed_model=np.array([EMBED_MODEL_NAME]),
            )
            tmp.flush()
            os.fsync(tmp.fileno())
        os.replace(temp_path, INDEX_DIR / "index.npz")
    finally:
        if temp_path and temp_path.exists():
            temp_path.unlink()


def validate_index(docs, embeddings) -> Dict:
    """손상되었거나 오래된 인덱스를 검색 단계까지 넘기지 않는다."""
    if not isinstance(docs, list) or not all(
        isinstance(doc, dict) and isinstance(doc.get("source"), str)
        and isinstance(doc.get("text"), str) for doc in docs
    ):
        raise ValueError("문서 목록 형식이 잘못되었습니다.")
    vectors = np.asarray(embeddings, dtype=np.float32)
    if vectors.ndim != 2 or vectors.shape[0] != len(docs) or vectors.shape[1] == 0:
        raise ValueError("문서 수와 임베딩 배열이 일치하지 않습니다.")
    return {"docs": docs, "embeddings": vectors}


def load_index() -> Optional[Dict]:
    """저장된 인덱스를 디스크에서 복원."""
    global _rag_engine, _rag_load_error
    archive = INDEX_DIR / "index.npz"
    ep, dp = INDEX_DIR / "embeddings.npy", INDEX_DIR / "docs.json"
    if not archive.exists() and not (ep.exists() or dp.exists()):
        _rag_load_error = None
        return None
    try:
        if archive.exists():
            with np.load(archive, allow_pickle=False) as data:
                saved_model = str(data["embed_model"][0])
                if saved_model != EMBED_MODEL_NAME:
                    raise ValueError(
                        f"索引使用 {saved_model}，当前嵌入模型为 {EMBED_MODEL_NAME}。"
                    )
                docs = json.loads(data["docs"].tobytes().decode("utf-8"))
                engine = validate_index(docs, data["embeddings"])
        else:
            if not (ep.exists() and dp.exists()):
                raise ValueError("旧索引文件不完整。")
            docs = json.loads(dp.read_text(encoding="utf-8"))
            engine = validate_index(docs, np.load(ep, allow_pickle=False))
        with _rag_lock:
            _rag_engine = engine
            _rag_load_error = None
        return engine
    except Exception as e:
        with _rag_lock:
            _rag_engine = None
            _rag_load_error = str(e)
        print(f"[RAG 로드 실패] {_rag_load_error}")
        return None


def remove_uploaded_file(name: str) -> int:
    """원본 파일과 그 파일의 검색 청크를 함께 제거한다."""
    global _rag_engine
    with _rag_lock:
        if _rag_load_error:
            raise RuntimeError(f"索引无法加载：{_rag_load_error}。请先备份或清空索引。")
        indexed = bool(_rag_engine and any(
            doc["source"] == name for doc in _rag_engine["docs"]
        ))
        try:
            validate_upload_name(name)
            path = UPLOADS_DIR / name
            stored = path.is_file()
        except ValueError:
            path, stored = None, False
        if not indexed and not stored:
            raise FileNotFoundError(name)

        engine = None
        if _rag_engine:
            keep = [i for i, doc in enumerate(_rag_engine["docs"])
                    if doc["source"] != name]
            engine = {
                "docs": [_rag_engine["docs"][i] for i in keep],
                "embeddings": _rag_engine["embeddings"][keep],
            }

        quarantine = None
        if stored:
            quarantine = UPLOADS_DIR / f".deleting-{uuid.uuid4().hex}"
            os.replace(path, quarantine)
        try:
            if engine is not None:
                save_index(engine)
        except Exception:
            if quarantine and quarantine.exists():
                os.replace(quarantine, path)
            raise
        if quarantine:
            quarantine.unlink()
        if engine is not None:
            _rag_engine = engine
        return len(engine["docs"]) if engine else 0


# ─── RAG 모드 결정 임계값 (multilingual-e5-base 정규화 코사인 기준) ───
# strict 이상이면 "자료 기반(grounded)" — 자료의 사실만 사용
# min 이하는 그 청크를 아예 참조하지 않음 (잡음 제거)
RAG_STRICT_SIM = 0.55
RAG_MIN_SIM = 0.40


def search_rag_per_doc(query: str, k_per_doc: int = 3):
    """[최적 다중문서 검색]
    각 업로드 파일(source)마다 top-k 청크를 뽑아 반환한다.
    - 글로벌 top-k가 한 파일에 쏠리는 쏠림 현상 방지
    - 서로 다른 문서를 LLM이 각각 참조할 수 있도록 그룹화

    Returns:
      (per_doc_list, max_similarity)
        per_doc_list = [{source, chunks:[{text, score}, ...]}, ...] — 최고 유사도 순
        max_similarity = 전역 최대 유사도 (grounded/fallback 모드 판정에 사용)
    """
    with _rag_lock:
        engine = _rag_engine
    if engine is None:
        return [], 0.0
    embedder = get_embedder()
    q_vec = embedder.encode([f"query: {query}"], normalize_embeddings=True)[0]
    sims = engine["embeddings"] @ q_vec
    docs = engine["docs"]
    if len(sims) == 0:
        return [], 0.0

    # 1) 소스별 청크 인덱스 그룹화
    by_source: Dict[str, List[int]] = {}
    for i, d in enumerate(docs):
        by_source.setdefault(d.get("source", "자료"), []).append(i)

    # 2) 각 소스에서 top-k 선택 (RAG_MIN_SIM 미만 제외)
    results = []
    for source, idxs in by_source.items():
        local = sorted(((i, float(sims[i])) for i in idxs), key=lambda t: -t[1])
        picked = []
        for i, sim in local[:k_per_doc]:
            if sim < RAG_MIN_SIM:
                break
            picked.append({"text": docs[i]["text"], "score": sim})
        if picked:
            results.append({"source": source, "chunks": picked})

    # 3) 각 소스를 "최고 청크 점수" 기준 내림차순 정렬
    results.sort(key=lambda r: -r["chunks"][0]["score"])
    return results, float(np.max(sims))


def decide_rag_mode(per_doc: List[Dict], max_sim: float) -> str:
    """'grounded' (자료 안에서만 답변) / 'fallback' (자료 범위 밖 · LLM 일반 지식)."""
    if not per_doc:
        return "fallback"
    return "grounded" if max_sim >= RAG_STRICT_SIM else "fallback"


def format_rag_context(per_doc: List[Dict]) -> str:
    """시스템 프롬프트에 삽입할 참고자료 블록.
    각 문서는 '[자료 N · 파일명]'으로 라벨링되어 모델이 출처별로 구분 참조 가능."""
    parts = []
    for i, d in enumerate(per_doc):
        chunks_txt = "\n···\n".join(c["text"] for c in d["chunks"])
        parts.append(f"[자료 {i+1} · {d['source']}]\n{chunks_txt}")
    return "\n\n".join(parts)


# ═══════════════════════════════════════════════════════
# 프롬프트 구성 (2-stage: grounded → fallback)
# ═══════════════════════════════════════════════════════
# 자료가 아예 없을 때 (업로드 전) 쓰이는 기본 프롬프트
SYSTEM_PROMPT_PLAIN = (
    "당신은 친절한 챗봇입니다. "
    "확실하지 않은 내용은 솔직하게 모른다고 말하세요. "
    "사용자의 마지막 질문과 같은 언어로 자연스럽고 명확하게 답변하세요."
)

# 자료가 질의와 잘 매칭될 때 (max_sim ≥ strict) — "자료 기반" 모드
SYSTEM_PROMPT_GROUNDED = (
    "당신은 친절한 챗봇입니다.\n\n"
    "[RAG 모드: 자료 기반 답변]\n"
    "- 아래 [참고 자료]는 사용자가 업로드한 서로 다른 문서에서 각각 발췌되었습니다 ('[자료 N · 파일명]').\n"
    "- 1차로는 자료 안의 사실만 사용해 답변하세요. 서로 다른 자료의 내용은 각각 명확히 구분해 참조하세요.\n"
    "- 답변 안에서 어느 자료를 참조했는지 '(자료 1)', '(자료 2)' 형태로 짧게 표기하세요.\n"
    "- 자료들끼리 상충하면 어느 쪽이 어떻게 다른지 간단히 짚으세요.\n"
    "- 질문이 자료 범위를 벗어나는 부분이 있으면: 그 부분에 한해 '(참고 자료 범위 밖 · 일반 지식)' 라고 명시한 뒤 LLM 일반 지식으로 보충하세요.\n"
    "- 사용자의 마지막 질문과 같은 언어로 답변하세요. 중국어 질문에는 중국어로, 한국어 질문에는 한국어로 답변하세요."
)

# 자료는 업로드되었지만 질의와 매칭이 약할 때 — "자료 범위 밖" 모드
SYSTEM_PROMPT_FALLBACK = (
    "당신은 친절한 챗봇입니다.\n\n"
    "[RAG 모드: 자료 범위 밖]\n"
    "- 업로드된 자료에서는 이 질문과 관련된 내용을 찾지 못했습니다.\n"
    "- 답변 서두에 업로드 자료 범위 밖의 일반 지식 기반 답변임을 사용자의 언어로 알리세요.\n"
    "- 확실하지 않은 내용은 솔직하게 모른다고 말하세요. 사용자의 마지막 질문과 같은 언어로 답변하세요."
)


def build_messages(user_input: str, history: List,
                   per_doc: Optional[List[Dict]] = None,
                   mode: str = "fallback",
                   model_id: Optional[str] = None,
                   language: Optional[Literal["zh", "ko"]] = None) -> List[Dict]:
    """LLM에 전달할 chat-completion 메시지 배열.
    - mode='grounded': 자료 기반 시스템 프롬프트 + 참고자료 블록
    - mode='fallback' + per_doc 있음: 자료는 있지만 매칭 약함 → "범위 밖" 안내 프롬프트
    - mode='fallback' + 자료 없음: 평범한 챗봇 프롬프트
    """
    if mode == "grounded" and per_doc:
        sys = SYSTEM_PROMPT_GROUNDED + f"\n\n[참고 자료]\n{format_rag_context(per_doc)}"
    elif _rag_engine and _rag_engine["docs"]:
        sys = SYSTEM_PROMPT_FALLBACK
    else:
        sys = SYSTEM_PROMPT_PLAIN

    if model_id:
        sys += (
            f"\n\n이 요청에서 선택한 모델 ID는 '{model_id}'입니다. "
            "모델을 물으면 이 설정된 ID를 정확히 알려주세요. "
            "제공업체의 내부 모델 버전은 확인하지 못했다면 단정하지 마세요."
        )

    if language == "zh":
        sys += "\n\n用户已明确选择中文。请用简体中文回答，即使提问或参考资料使用其他语言；此选择优先于上面的语言规则。"
    elif language == "ko":
        sys += "\n\n사용자가 한국어를 명시적으로 선택했습니다. 질문이나 참고 자료의 언어와 관계없이 한국어로 답변하세요. 이 선택은 위의 언어 규칙보다 우선합니다."

    msgs = [{"role": "system", "content": sys}]
    # 최근 대화 6개만 포함 (컨텍스트 길이 절약)
    for h in history[-6:]:
        if isinstance(h, list) and len(h) == 2 and h[1]:
            role = h[0] if h[0] in ("user", "assistant") else "user"
            msgs.append({"role": role, "content": h[1]})
    msgs.append({"role": "user", "content": user_input})
    return msgs


# ═══════════════════════════════════════════════════════
# FastAPI
# ═══════════════════════════════════════════════════════
@asynccontextmanager
async def lifespan(_app: FastAPI):
    remote = HOST.lower() not in ("localhost", "127.0.0.1", "::1")
    if remote and not APP_ACCESS_TOKEN:
        raise ValueError("对外监听时必须设置 APP_ACCESS_TOKEN，避免公开付费 API 和资料接口。")
    if remote and len(APP_ACCESS_TOKEN) < 16:
        raise ValueError("对外监听时 APP_ACCESS_TOKEN 至少需要 16 个字符。")
    catalog = get_model_catalog()
    api_count = sum(model["provider"] == "api" for model in catalog.values())
    if api_count:
        print(f"[startup] API 모델 {api_count}개 등록됨")
    eng = load_index()
    if eng:
        print(f"[startup] RAG 인덱스 자동 로드: {len(eng['docs'])}개 청크")
    default_key = get_default_model_key(catalog)
    if default_key and catalog[default_key]["provider"] == "local":
        print(f"[startup] 기본 로컬 모델: {default_key}. 백그라운드 로딩 시작.")
        threading.Thread(target=lambda: get_llm(default_key), daemon=True).start()
    elif default_key is None:
        print("\n[안내] 사용 가능한 모델이 없습니다. 로컬 모델 파일·llama-cpp 설치 또는 서버 API Key를 확인하세요.\n")
    yield


app = FastAPI(title="FICL Chatbot", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[f"http://localhost:{PORT}", f"http://127.0.0.1:{PORT}"],
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["Content-Type"],
)


def require_access(request: Request) -> None:
    """로컬은 기본 허용, 토큰이 설정된 서버는 모든 데이터 API를 보호한다."""
    if not APP_ACCESS_TOKEN:
        return
    scheme, _, credential = request.headers.get("Authorization", "").partition(" ")
    if scheme.lower() != "bearer" or not hmac.compare_digest(
        credential, APP_ACCESS_TOKEN
    ):
        raise HTTPException(401, "访问令牌无效或未提供。")


class ChatRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    message: str
    history: List = Field(default_factory=list)
    model: Optional[str] = None
    language: Optional[Literal["zh", "ko"]] = None


class VerifyModelRequest(BaseModel):
    model: str


@app.get("/")
async def root():
    if UI_HTML.exists():
        return HTMLResponse(UI_HTML.read_text(encoding="utf-8"))
    return HTMLResponse("<h1>ui.html 파일이 없습니다.</h1>")


@app.get("/guide")
async def guide():
    if GUIDE_HTML.exists():
        return HTMLResponse(GUIDE_HTML.read_text(encoding="utf-8"))
    raise HTTPException(404, "guide.html 文件不存在。")


@app.get("/status", dependencies=[Depends(require_access)])
async def status():
    """모델 목록·다운로드 여부·RAG 상태 반환."""
    catalog = get_model_catalog()
    return {
        "models": [public_model_info(model) for model in catalog.values()],
        "default_model": get_default_model_key(catalog),
        "rag_chunks": len(_rag_engine["docs"]) if _rag_engine else 0,
        "rag_error": _rag_load_error,
        "llama_ok": LLAMA_OK,
        "embedder_ok": ST_OK,
    }


@app.post("/verify_model", dependencies=[Depends(require_access)])
async def verify_model(req: VerifyModelRequest):
    """사용자 버튼으로만 호출한다. 소량의 외부 API 사용량이 발생할 수 있다."""
    selected = select_model(req.model, get_model_catalog())
    if selected["provider"] != "api":
        raise HTTPException(400, "仅 API 模型需要连接测试。")
    try:
        await run_in_threadpool(verify_api_model, selected)
    except Exception as e:
        detail = str(e)
        secret = selected.get("api_key", "")
        if secret:
            detail = detail.replace(secret, "***")
        raise HTTPException(502, detail) from e
    return {"success": True, "model": selected["key"]}


@app.get("/files", dependencies=[Depends(require_access)])
async def files_list():
    return {"files": await run_in_threadpool(list_uploaded_files)}


@app.delete("/files", dependencies=[Depends(require_access)])
async def files_delete(name: str = Query(..., min_length=1)):
    try:
        chunks = await run_in_threadpool(remove_uploaded_file, name)
    except FileNotFoundError:
        raise HTTPException(404, "文件不存在。") from None
    except RuntimeError as e:
        raise HTTPException(409, str(e)) from e
    return {"success": True, "name": name, "chunks": chunks}


@app.post("/upload", dependencies=[Depends(require_access)])
async def upload(files: List[UploadFile] = File(...)):
    """파일 업로드 → 텍스트 추출 → 임베딩 → 인덱스 저장."""
    if len(files) > MAX_UPLOAD_FILES:
        raise HTTPException(413, f"一次最多上传 {MAX_UPLOAD_FILES} 个文件。")
    file_list = []
    total_bytes = 0
    seen_names = set()
    for f in files:
        name = (f.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
        try:
            validate_upload_name(name)
        except ValueError as e:
            raise HTTPException(400, str(e)) from e
        if name.casefold() in seen_names:
            raise HTTPException(400, f"本次上传有同名文件：{name}")
        seen_names.add(name.casefold())
        data = await f.read(MAX_UPLOAD_BYTES + 1)
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(413, f"文件超过 20 MB：{name}")
        total_bytes += len(data)
        if total_bytes > MAX_TOTAL_UPLOAD_BYTES:
            raise HTTPException(413, "本次上传文件总量不得超过 50 MB。")
        file_list.append({"name": name, "data": data})

    try:
        engine = await run_in_threadpool(build_index, file_list)
    except (ValueError, RuntimeError) as e:
        raise HTTPException(400, str(e)) from e
    if not engine:
        raise HTTPException(400, "未能从文件中提取可索引的文字。")

    return {"success": True, "chunks": len(engine["docs"])}


@app.post("/reset_index", dependencies=[Depends(require_access)])
async def reset_index():
    """원본 업로드와 RAG 인덱스를 모두 초기화한다."""
    global _rag_engine, _rag_load_error
    with _rag_lock:
        for path in UPLOADS_DIR.iterdir():
            if path.is_file():
                path.unlink()
        for fn in ("index.npz", "embeddings.npy", "docs.json"):
            p = INDEX_DIR / fn
            if p.exists():
                p.unlink()
        _rag_engine = None
        _rag_load_error = None
    return {"success": True}


@app.post("/chat", dependencies=[Depends(require_access)])
async def chat(req: ChatRequest):
    """채팅 — SSE로 토큰 단위 스트리밍.
    흐름:
      1) 각 업로드 문서별 top-k 청크 수집 (search_rag_per_doc)
      2) 전역 최대 유사도로 grounded/fallback 모드 결정
      3) 모드별 시스템 프롬프트 구성 (build_messages)
      4) meta 이벤트 → 토큰 스트림 → done 이벤트 순으로 전송
    """
    user_input = req.message.strip()
    if not user_input:
        raise HTTPException(400, "消息不能为空。")
    selected = select_model(req.model, get_model_catalog())
    selected_key = selected["key"]
    provider = selected["provider"]
    selected_api = selected if provider == "api" else None

    # ── 1단계: 다중문서 RAG 검색 + 모드 판정 ─────────────
    per_doc, max_sim = await run_in_threadpool(search_rag_per_doc, user_input, 3)
    rag_mode = decide_rag_mode(per_doc, max_sim)
    rag_available = bool(_rag_engine and _rag_engine["docs"])
    rag_sources = [d["source"] for d in per_doc]

    model_id = selected_api["model_id"] if selected_api else selected_key
    msgs = build_messages(user_input, req.history, per_doc=per_doc,
                          mode=rag_mode, model_id=model_id, language=req.language)

    def event_stream():
        # ── 스트림 시작 전: 메타 정보 먼저 전송 (UI가 뱃지 표시 가능) ─
        meta = {
            "meta": True,
            "rag_available": rag_available,
            "rag_mode": rag_mode,
            "rag_sources": rag_sources,
            "rag_max_sim": round(max_sim, 3),
            "provider": provider,
            "model_key": selected_key,
            "model_name": selected["name"],
        }
        yield f"data: {json.dumps(meta, ensure_ascii=False)}\n\n"

        try:
            full = ""
            if provider == "api":
                token_stream = stream_api_completion(msgs, selected_api)
                for token in token_stream:
                    full += token
                    yield f"data: {json.dumps({'token': token}, ensure_ascii=False)}\n\n"
            else:
                llm = get_llm(selected_key)
                stream = llm.create_chat_completion(
                    messages=msgs,
                    max_tokens=1024,
                    temperature=0.7,
                    top_p=0.9,
                    stream=True,
                )
                for chunk in stream:
                    # 첫 청크는 role만, 마지막 청크는 finish_reason만 있을 수 있음
                    try:
                        delta = chunk.get("choices", [{}])[0].get("delta", {}).get("content", "")
                    except (KeyError, IndexError, AttributeError):
                        delta = ""
                    if delta:
                        full += delta
                        yield f"data: {json.dumps({'token': delta}, ensure_ascii=False)}\n\n"
            payload = {
                "done": True,
                "text": full,
                "sources": rag_sources,     # 호환성 유지 (기존 UI 필드)
                "rag_mode": rag_mode,
                "rag_max_sim": round(max_sim, 3),
                "provider": provider,
                "model_key": selected_key,
                "model_name": selected["name"],
            }
            yield f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"
        except FileNotFoundError as e:
            yield f"data: {json.dumps({'error': str(e)}, ensure_ascii=False)}\n\n"
        except Exception as e:
            detail = str(e)
            if selected_api is not None:
                secret = selected_api.get("api_key", "")
                if secret:
                    detail = detail.replace(secret, "***")
            yield f"data: {json.dumps({'error': detail}, ensure_ascii=False)}\n\n"

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


if __name__ == "__main__":
    import uvicorn
    print(f"\n{'='*50}")
    print(f"  FICL Chatbot")
    print(f"  http://{HOST}:{PORT}")
    print(f"{'='*50}\n")
    uvicorn.run("chatbot:app", host=HOST, port=PORT, reload=False)
